import uuid
from datetime import date, datetime

from fastapi import APIRouter, Depends, File, HTTPException, Request, UploadFile
from openpyxl import Workbook, load_workbook
from pydantic import BaseModel, ConfigDict, Field
from sqlalchemy.orm import Session

from app.api.audit_log import write_audit_log_entry
from app.core.auth import require_roles
from app.db.session import get_db
from app.models.setting import DocumentType, Override, Setting, SettingScope
from app.models.user import User
from app.xlsx_utils import xlsx_header_row, xlsx_response

settings_router = APIRouter(prefix="/settings", tags=["settings"])
overrides_router = APIRouter(prefix="/overrides", tags=["overrides"])

# Q.2 rule 6: "Master Settings screen is Director-only (PM read-only)."
READ_ROLES = ("pm", "director")
WRITE_ROLES = ("director",)
# Overrides are logged by whoever is working the document -- PM/Director,
# consistent with every other cost-adjacent write in this build (K.3).
OVERRIDE_ROLES = ("pm", "director")


def get_current_setting_value(db: Session, key: str, scope_value: str | None = None) -> str | None:
    """The most recent version of `key` whose effective_from has arrived.
    Q.2 rule 1: editing a setting inserts a new version rather than
    mutating the old one, so this never changes what an already-frozen
    document read."""
    today = date.today()
    row = (
        db.query(Setting)
        .filter(Setting.key == key, Setting.scope_value == scope_value, Setting.effective_from <= today)
        .order_by(Setting.effective_from.desc(), Setting.created_at.desc())
        .first()
    )
    return row.value if row else None


def get_gst_rate_percent(db: Session, default: float = 18.0) -> float:
    """K.4: 'The GST rate itself (18%) is a Master Setting (Q.1), not
    hard-coded, in case it ever changes.' Falls back to the blueprint's
    own default when no Setting row exists yet (e.g. a fresh test DB)."""
    value = get_current_setting_value(db, "gst_rate_percent")
    return float(value) if value is not None else default


def get_internal_email_domains(db: Session) -> set[str]:
    """Q.1 Communications 'internal-domain list' -- the domains Cost
    Sheet/consumption-sheet/BOM messages (M.7.2 rule 7) may be emailed
    to. Unlike gst_rate_percent, there's no sensible universal default
    for a company's own email domain, so until the Director sets
    `internal_email_domains` explicitly (comma-separated) this falls
    back to the domains already in use by this installation's own USERS
    -- self-configuring from real data instead of a fabricated default,
    and still fail-closed (a domain nobody here has ever logged in from
    is never treated as internal)."""
    value = get_current_setting_value(db, "internal_email_domains")
    if value:
        return {domain.strip().lower() for domain in value.split(",") if domain.strip()}
    return {
        email.split("@", 1)[1].lower()
        for (email,) in db.query(User.email).all()
        if "@" in email
    }


def _current_settings(db: Session, key_prefix: str | None = None) -> list[Setting]:
    """One row per (key, scope_value): its currently-effective version.
    Shared by the GET list, bulk %, and Excel export endpoints (Q.2 rule 6)."""
    today = date.today()
    query = db.query(Setting).filter(Setting.effective_from <= today)
    if key_prefix is not None:
        query = query.filter(Setting.key.like(f"{key_prefix}%"))
    all_rows = query.order_by(
        Setting.key, Setting.scope_value, Setting.effective_from.desc(), Setting.created_at.desc()
    ).all()
    seen: set[tuple[str, str | None]] = set()
    current: list[Setting] = []
    for row in all_rows:
        identity = (row.key, row.scope_value)
        if identity in seen:
            continue
        seen.add(identity)
        current.append(row)
    return current


class SettingOut(BaseModel):
    id: uuid.UUID
    key: str
    scope: SettingScope
    scope_value: str | None
    value: str
    unit: str | None
    effective_from: date
    changed_by_id: uuid.UUID
    reason: str | None
    created_at: datetime

    model_config = ConfigDict(from_attributes=True)


@settings_router.get("", response_model=list[SettingOut])
def list_current_settings(
    db: Session = Depends(get_db),
    current_user=Depends(require_roles(*READ_ROLES)),
):
    return _current_settings(db)


@settings_router.get("/{key}/history", response_model=list[SettingOut])
def get_setting_history(
    key: str,
    scope_value: str | None = None,
    db: Session = Depends(get_db),
    current_user=Depends(require_roles(*READ_ROLES)),
):
    rows = (
        db.query(Setting)
        .filter(Setting.key == key, Setting.scope_value == scope_value)
        .order_by(Setting.effective_from.desc(), Setting.created_at.desc())
        .all()
    )
    if not rows:
        raise HTTPException(status_code=404, detail="No such setting")
    return rows


class SettingCreate(BaseModel):
    key: str
    scope: SettingScope = SettingScope.GLOBAL
    scope_value: str | None = None
    value: str
    unit: str | None = None
    effective_from: date | None = None
    reason: str | None = None


@settings_router.post("", response_model=SettingOut, status_code=201)
def create_setting_version(
    payload: SettingCreate,
    request: Request,
    db: Session = Depends(get_db),
    current_user=Depends(require_roles(*WRITE_ROLES)),
):
    """Q.2 rule 1: 'editing' a Master Setting is really creating a new
    version, effective from a given date -- Director only."""
    old_value = get_current_setting_value(db, payload.key, payload.scope_value)
    setting = Setting(
        key=payload.key,
        scope=payload.scope,
        scope_value=payload.scope_value,
        value=payload.value,
        unit=payload.unit,
        effective_from=payload.effective_from or date.today(),
        changed_by_id=current_user.id,
        reason=payload.reason,
    )
    db.add(setting)
    db.flush()

    # Q.2 rule 7: "Settings changes are in the audit log (M.5)."
    write_audit_log_entry(
        db, current_user, "setting", setting.id, payload.key,
        old_value=old_value, new_value=payload.value, reason=payload.reason, request=request,
    )

    db.commit()
    db.refresh(setting)
    return setting


class BulkUpdateRequest(BaseModel):
    key_prefix: str
    percent_change: float  # e.g. 6.0 for "+6%", -5.0 for "-5%"
    effective_from: date | None = None
    reason: str


@settings_router.post("/bulk-update", response_model=list[SettingOut])
def bulk_update_settings(
    payload: BulkUpdateRequest,
    request: Request,
    db: Session = Depends(get_db),
    current_user=Depends(require_roles(*WRITE_ROLES)),
):
    """Q.2 rule 5: 'Director can apply a % change to a whole category
    (e.g. "steel +6%") with one action and an effective date.' A
    'category' here is every setting whose key starts with key_prefix;
    each gets its own new version, same effective_from and reason."""
    current = _current_settings(db, key_prefix=payload.key_prefix)

    if not current:
        raise HTTPException(status_code=404, detail=f"No settings found with key prefix '{payload.key_prefix}'")

    effective_from = payload.effective_from or date.today()
    new_versions = []
    for row in current:
        try:
            new_value = float(row.value) * (1 + payload.percent_change / 100)
        except ValueError:
            continue  # non-numeric setting values aren't bulk-adjustable
        new_setting = Setting(
            key=row.key,
            scope=row.scope,
            scope_value=row.scope_value,
            value=str(round(new_value, 4)),
            unit=row.unit,
            effective_from=effective_from,
            changed_by_id=current_user.id,
            reason=payload.reason,
        )
        db.add(new_setting)
        db.flush()
        new_versions.append(new_setting)

        # Q.2 rule 7: "Settings changes are in the audit log (M.5)."
        write_audit_log_entry(
            db, current_user, "setting", new_setting.id, row.key,
            old_value=row.value, new_value=new_setting.value, reason=payload.reason, request=request,
        )

    db.commit()
    for v in new_versions:
        db.refresh(v)
    return new_versions


# --------------------------------------------------------------------------
# Excel export/import (Q.2 rule 6: "Master Settings screen is Director-only
# (PM read-only), exportable to Excel and importable back, so NestaPrime
# can maintain its rate card in Excel if preferred and upload it.")
# --------------------------------------------------------------------------

_XLSX_COLUMNS = ["Key", "Scope", "Scope value", "Value", "Unit", "Effective from", "Reason"]
_IMPORT_DEFAULT_REASON = "Bulk Excel import"


@settings_router.get("/export")
def export_settings(
    db: Session = Depends(get_db),
    current_user=Depends(require_roles(*READ_ROLES)),
):
    """One row per currently-effective setting -- the exact same set
    GET /settings returns, as a downloadable workbook. Re-uploading this
    file unchanged via POST /settings/import creates nothing (see that
    endpoint), so export-then-reimport is always safe to run."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Master Settings"
    xlsx_header_row(ws, _XLSX_COLUMNS)
    for row in _current_settings(db):
        ws.append([row.key, row.scope.value, row.scope_value, row.value, row.unit, row.effective_from.isoformat(), row.reason])
    return xlsx_response(wb, "master-settings.xlsx")


class SettingImportRowError(BaseModel):
    row: int
    detail: str


class SettingImportResult(BaseModel):
    created: list[SettingOut]
    unchanged: int
    errors: list[SettingImportRowError]


def _parse_import_effective_from(raw) -> date:
    if raw is None or raw == "":
        return date.today()
    if isinstance(raw, datetime):
        return raw.date()
    if isinstance(raw, date):
        return raw
    return date.fromisoformat(str(raw).strip())


@settings_router.post("/import", response_model=SettingImportResult)
def import_settings(
    request: Request,
    file: UploadFile = File(...),
    db: Session = Depends(get_db),
    current_user=Depends(require_roles(*WRITE_ROLES)),
):
    """Q.2 rule 6's 'importable back.' Expects the same columns
    export_settings() produces (Key, Scope, Scope value, Value, Unit,
    Effective from, Reason) in that order on the first worksheet, header
    row first. Each data row becomes a new Setting VERSION (Q.2 rule 1 --
    never a mutation) unless its Value is identical to what's already
    currently effective for that (key, scope_value), in which case it's
    counted as unchanged and skipped -- so re-exporting and reimporting
    the same file is a safe no-op, and only the rows a Director actually
    edited in Excel produce new versions. One bad row (unknown scope, an
    unparseable date, a missing key/value) is recorded as a per-row error
    and does not stop the rest of the file from importing."""
    try:
        wb = load_workbook(file.file, data_only=True)
    except Exception as exc:
        raise HTTPException(status_code=422, detail=f"Could not read this file as an Excel workbook: {exc}")
    ws = wb.active

    valid_scopes = {s.value for s in SettingScope}
    created: list[Setting] = []
    unchanged = 0
    errors: list[SettingImportRowError] = []

    for row_number, row in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
        if row is None or all(cell in (None, "") for cell in row):
            continue  # blank row -- e.g. Excel's own trailing rows
        key_cell, scope_cell, scope_value_cell, value_cell, unit_cell, effective_from_cell, reason_cell = (
            list(row) + [None] * (len(_XLSX_COLUMNS) - len(row))
        )[: len(_XLSX_COLUMNS)]

        try:
            key = str(key_cell).strip() if key_cell not in (None, "") else ""
            if not key:
                raise ValueError("Key is required")
            value = str(value_cell).strip() if value_cell not in (None, "") else ""
            if not value:
                raise ValueError("Value is required")
            scope_raw = str(scope_cell).strip().lower() if scope_cell not in (None, "") else SettingScope.GLOBAL.value
            if scope_raw not in valid_scopes:
                raise ValueError(f"Unknown scope '{scope_raw}' -- must be one of {sorted(valid_scopes)}")
            scope_value = str(scope_value_cell).strip() if scope_value_cell not in (None, "") else None
            unit = str(unit_cell).strip() if unit_cell not in (None, "") else None
            effective_from = _parse_import_effective_from(effective_from_cell)
            reason = str(reason_cell).strip() if reason_cell not in (None, "") else _IMPORT_DEFAULT_REASON
        except (ValueError, TypeError) as exc:
            errors.append(SettingImportRowError(row=row_number, detail=str(exc)))
            continue

        old_value = get_current_setting_value(db, key, scope_value)
        if old_value == value:
            unchanged += 1
            continue

        setting = Setting(
            key=key,
            scope=SettingScope(scope_raw),
            scope_value=scope_value,
            value=value,
            unit=unit,
            effective_from=effective_from,
            changed_by_id=current_user.id,
            reason=reason,
        )
        db.add(setting)
        db.flush()
        write_audit_log_entry(
            db, current_user, "setting", setting.id, key,
            old_value=old_value, new_value=value, reason=reason, request=request,
        )
        created.append(setting)

    db.commit()
    for s in created:
        db.refresh(s)
    return SettingImportResult(
        created=[SettingOut.model_validate(s) for s in created], unchanged=unchanged, errors=errors
    )


# --------------------------------------------------------------------------
# Overrides (Part O OVERRIDES / Q.2 rule 2)
# --------------------------------------------------------------------------


class OverrideCreate(BaseModel):
    document_type: DocumentType
    document_id: uuid.UUID
    setting_key: str
    master_value: str
    override_value: str
    reason: str = Field(min_length=1)


class OverrideOut(BaseModel):
    id: uuid.UUID
    document_type: DocumentType
    document_id: uuid.UUID
    setting_key: str
    master_value: str
    override_value: str
    reason: str
    user_id: uuid.UUID
    created_at: datetime

    model_config = ConfigDict(from_attributes=True)


@overrides_router.post("", response_model=OverrideOut, status_code=201)
def create_override(
    payload: OverrideCreate,
    db: Session = Depends(get_db),
    current_user=Depends(require_roles(*OVERRIDE_ROLES)),
):
    override = Override(
        document_type=payload.document_type,
        document_id=payload.document_id,
        setting_key=payload.setting_key,
        master_value=payload.master_value,
        override_value=payload.override_value,
        reason=payload.reason,
        user_id=current_user.id,
    )
    db.add(override)
    db.commit()
    db.refresh(override)
    return override


@overrides_router.get("", response_model=list[OverrideOut])
def list_overrides(
    document_type: DocumentType,
    document_id: uuid.UUID,
    db: Session = Depends(get_db),
    current_user=Depends(require_roles(*OVERRIDE_ROLES)),
):
    return (
        db.query(Override)
        .filter(Override.document_type == document_type, Override.document_id == document_id)
        .order_by(Override.created_at.desc())
        .all()
    )
