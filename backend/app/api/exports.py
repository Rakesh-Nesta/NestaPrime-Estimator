import io
import uuid

from fastapi import APIRouter, Depends, HTTPException
from fastapi.responses import StreamingResponse
from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter
from sqlalchemy.orm import Session

from app.api import schedule as schedule_api
from app.api.documents import _po_lookup_for_cost_sheet
from app.core.auth import require_roles
from app.db.session import get_db
from app.models.client import Client
from app.models.document import (
    CostSheet,
    CostSheetLine,
    EstimateOption,
    Quotation,
    QuotationLine,
    QuotationStatus,
)
from app.models.project import Project
from app.models.sport import ProjectSport, Sport

exports_router = APIRouter(tags=["exports"])

# J.4: Cost Sheet / Consumption Sheet / BOM are all "internal only" -- same
# cost-visibility gate as the documents they're exported from (K.3).
COST_ROLES = ("pm", "director")
# The Vendor RFQ is the one export explicitly meant to leave the building
# (sent to vendors) -- Procurement needs it even though Procurement has no
# access to the cost-bearing exports above.
RFQ_ROLES = ("pm", "director", "procurement")
# Billing Handoff carries no cost/margin (K.3-safe by construction), so
# Sales -- who creates and sends the Quotation -- can use it too.
HANDOFF_ROLES = ("sales", "pm", "director")


def _xlsx_response(wb: Workbook, filename: str) -> StreamingResponse:
    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return StreamingResponse(
        buffer,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="{filename}"'},
    )


def _header_row(ws, headers: list[str]) -> None:
    ws.append(headers)
    for cell in ws[1]:
        cell.font = Font(bold=True)
    for i, header in enumerate(headers, start=1):
        ws.column_dimensions[get_column_letter(i)].width = max(12, len(header) + 2)


def _get_cost_sheet_or_404(db: Session, cost_sheet_id: uuid.UUID) -> CostSheet:
    cost_sheet = db.query(CostSheet).filter(CostSheet.id == cost_sheet_id).first()
    if not cost_sheet:
        raise HTTPException(status_code=404, detail="Cost sheet not found")
    return cost_sheet


# --------------------------------------------------------------------------
# J.4: Cost Sheet (Excel, editable) -- "all lines, formulas live, cost
# price, labour, overheads; internal only."
# --------------------------------------------------------------------------


@exports_router.get("/cost-sheets/{cost_sheet_id}/exports/cost-sheet")
def export_cost_sheet(
    cost_sheet_id: uuid.UUID,
    db: Session = Depends(get_db),
    current_user=Depends(require_roles(*COST_ROLES)),
):
    cost_sheet = _get_cost_sheet_or_404(db, cost_sheet_id)
    lines = db.query(CostSheetLine).filter(CostSheetLine.cost_sheet_id == cost_sheet_id).all()

    wb = Workbook()
    ws = wb.active
    ws.title = "Cost Sheet"
    _header_row(
        ws,
        ["Work package", "Category", "Item", "Spec", "Unit", "Quantity", "Rate", "Amount", "Source", "Wastage %"],
    )
    for line in lines:
        row = ws.max_row + 1
        ws.append(
            [
                line.work_package.value,
                line.category,
                line.item_name,
                line.spec,
                line.unit,
                float(line.quantity),
                float(line.rate),
                None,  # Amount -- filled in as a live formula below
                line.source.value,
                float(line.wastage_percent) if line.wastage_percent is not None else None,
            ]
        )
        ws.cell(row=row, column=8, value=f"=F{row}*G{row}")  # Amount = Quantity x Rate, live

    last_row = ws.max_row
    total_row = last_row + 2
    ws.cell(row=total_row, column=7, value="Material + labour sum (this sheet):").font = Font(bold=True)
    if last_row >= 2:
        ws.cell(row=total_row, column=8, value=f"=SUM(H2:H{last_row})")
    ws.cell(row=total_row + 1, column=7, value="Cost incl. contingency (K.1, from /recompute):").font = Font(bold=True)
    ws.cell(row=total_row + 1, column=8, value=float(cost_sheet.cost_total))
    ws.cell(row=total_row + 2, column=7, value="(K.1's site-establishment %/contingency-by-package chain is computed").font = Font(italic=True, size=9)
    ws.cell(row=total_row + 3, column=7, value="by the app, not reproduced as spreadsheet formulas -- see /recompute.)").font = Font(italic=True, size=9)

    return _xlsx_response(wb, f"{cost_sheet.document_no}-cost-sheet.xlsx")


# --------------------------------------------------------------------------
# J.4: Material Consumption Sheet (Excel, editable) -- J.3's own columns.
# --------------------------------------------------------------------------


@exports_router.get("/cost-sheets/{cost_sheet_id}/exports/consumption-sheet")
def export_consumption_sheet(
    cost_sheet_id: uuid.UUID,
    db: Session = Depends(get_db),
    current_user=Depends(require_roles(*COST_ROLES)),
):
    cost_sheet = _get_cost_sheet_or_404(db, cost_sheet_id)
    lines = db.query(CostSheetLine).filter(CostSheetLine.cost_sheet_id == cost_sheet_id).all()
    po_by_cost_sheet_line_id = _po_lookup_for_cost_sheet(db, cost_sheet_id)

    wb = Workbook()
    ws = wb.active
    ws.title = "Consumption Sheet"
    _header_row(
        ws,
        [
            "Category", "Item & spec", "Unit", "Theoretical qty", "Wastage %", "Order qty",
            "Rate (internal)", "Amount (internal)", "Vendor", "PO No.", "Delivery date", "Received qty", "Balance",
        ],
    )
    for line in lines:
        row = ws.max_row + 1
        item_and_spec = line.item_name + (f" ({line.spec})" if line.spec else "")
        wastage = float(line.wastage_percent) if line.wastage_percent is not None else 0.0
        procurement = po_by_cost_sheet_line_id.get(line.id)
        po_line, po, vendor = procurement if procurement else (None, None, None)
        ws.append(
            [
                line.category,
                item_and_spec,
                line.unit,
                None,  # Theoretical qty -- live formula below
                wastage,
                float(line.quantity),  # Order qty is the stored, authoritative value
                float(line.rate),
                None,  # Amount -- live formula below
                vendor.name if vendor else None,
                po.po_no if po else None,
                po.delivery_date.date().isoformat() if po and po.delivery_date else None,
                float(po_line.received_qty) if po_line else None,
                (float(po_line.quantity) - float(po_line.received_qty)) if po_line else None,
            ]
        )
        ws.cell(row=row, column=4, value=f"=F{row}/(1+E{row}/100)")  # Theoretical = Order / (1 + wastage%)
        ws.cell(row=row, column=8, value=f"=F{row}*G{row}")  # Amount = Order qty x Rate

    last_row = ws.max_row
    note_row = last_row + 2
    ws.cell(row=note_row, column=1, value="Vendor / PO No. / Delivery date / Received qty / Balance are filled in "
                                           "once a Purchase Order (Part O) is raised for that line -- blank until "
                                           "then, not fabricated.").font = Font(italic=True, size=9)

    return _xlsx_response(wb, f"{cost_sheet.document_no}-consumption-sheet.xlsx")


# --------------------------------------------------------------------------
# J.4: Vendor RFQ (Excel/PDF) -- "item, spec, qty, unit only; no rates, no
# vendor names." Excel chosen over PDF (the blueprint offers either);
# Procurement gets access here despite having none of the cost-bearing
# exports, since this is the one meant to leave the building.
# --------------------------------------------------------------------------


@exports_router.get("/cost-sheets/{cost_sheet_id}/exports/rfq")
def export_rfq(
    cost_sheet_id: uuid.UUID,
    db: Session = Depends(get_db),
    current_user=Depends(require_roles(*RFQ_ROLES)),
):
    cost_sheet = _get_cost_sheet_or_404(db, cost_sheet_id)
    lines = db.query(CostSheetLine).filter(CostSheetLine.cost_sheet_id == cost_sheet_id).all()

    wb = Workbook()
    ws = wb.active
    ws.title = "Vendor RFQ"
    _header_row(ws, ["Item", "Spec", "Unit", "Qty"])
    for line in lines:
        ws.append([line.item_name, line.spec, line.unit, float(line.quantity)])

    return _xlsx_response(wb, f"{cost_sheet.document_no}-rfq.xlsx")


# --------------------------------------------------------------------------
# J.4: BOM (Excel) -- "internal, categorised, with source AI/Manual."
# --------------------------------------------------------------------------


@exports_router.get("/cost-sheets/{cost_sheet_id}/exports/bom")
def export_bom(
    cost_sheet_id: uuid.UUID,
    db: Session = Depends(get_db),
    current_user=Depends(require_roles(*COST_ROLES)),
):
    cost_sheet = _get_cost_sheet_or_404(db, cost_sheet_id)
    lines = (
        db.query(CostSheetLine)
        .filter(CostSheetLine.cost_sheet_id == cost_sheet_id)
        .order_by(CostSheetLine.category)  # "categorised"
        .all()
    )

    wb = Workbook()
    ws = wb.active
    ws.title = "BOM"
    _header_row(ws, ["Category", "Item", "Spec", "Unit", "Quantity", "Rate", "Amount", "Source"])
    for line in lines:
        row = ws.max_row + 1
        ws.append(
            [
                line.category, line.item_name, line.spec, line.unit,
                float(line.quantity), float(line.rate), None, line.source.value,
            ]
        )
        ws.cell(row=row, column=7, value=f"=E{row}*F{row}")

    return _xlsx_response(wb, f"{cost_sheet.document_no}-bom.xlsx")


# --------------------------------------------------------------------------
# J.4: Billing Handoff -- appears once a Quotation reaches Won. "Client
# name, contact, address, quotation number and date, the payment schedule/
# milestones, and the total (GST-inclusive) -- nothing else, no cost, no
# margin, no HSN." Manual handoff: NestaPrime enters the actual bill in
# Tally themselves -- this is a formatted read of data the Quotation and
# Schedule already hold, no new entity, no Tally integration.
# --------------------------------------------------------------------------


@exports_router.get("/quotations/{quotation_id}/exports/billing-handoff")
def export_billing_handoff(
    quotation_id: uuid.UUID,
    db: Session = Depends(get_db),
    current_user=Depends(require_roles(*HANDOFF_ROLES)),
):
    quotation = db.query(Quotation).filter(Quotation.id == quotation_id).first()
    if not quotation:
        raise HTTPException(status_code=404, detail="Quotation not found")
    if quotation.status != QuotationStatus.WON:
        raise HTTPException(
            status_code=400, detail="Billing Handoff is only available once a Quotation reaches Won"
        )

    project = db.query(Project).filter(Project.id == quotation.project_id).first()
    client = db.query(Client).filter(Client.id == project.client_id).first()

    wb = Workbook()
    ws = wb.active
    ws.title = "Billing Handoff"
    bold = Font(bold=True)

    info_rows = [
        ("Client", client.name),
        ("Contact", client.contact_name),
        ("Phone", client.phone),
        ("Email", client.email),
        ("Billing address", client.billing_address),
        ("Quotation No.", quotation.document_no),
        ("Date", quotation.sent_at.date().isoformat() if quotation.sent_at else quotation.created_at.date().isoformat()),
        ("Total (GST-inclusive, 18% flat)", float(quotation.quotation_total)),
    ]
    for label, value in info_rows:
        ws.append([label, value])
        ws.cell(row=ws.max_row, column=1).font = bold
    ws.column_dimensions["A"].width = 30
    ws.column_dimensions["B"].width = 40

    # Payment schedule per sport covered by the quotation (M.2 rule 11:
    # "covers the approved subset of sports") -- each sport's Part N
    # schedule is listed separately rather than a fabricated merge.
    ws.append([])
    ws.append(["Payment schedule"])
    ws.cell(row=ws.max_row, column=1).font = bold

    lines = db.query(QuotationLine).filter(QuotationLine.quotation_id == quotation_id).all()
    for qline in lines:
        option = db.query(EstimateOption).filter(EstimateOption.id == qline.estimate_option_id).first()
        project_sport = db.query(ProjectSport).filter(ProjectSport.id == option.project_sport_id).first()
        sport = db.query(Sport).filter(Sport.id == project_sport.sport_id).first()

        ws.append([sport.name])
        ws.cell(row=ws.max_row, column=1).font = Font(italic=True)
        header_row = ws.max_row + 1
        ws.append(["Milestone", "%", "Date"])
        for cell in ws[header_row]:
            cell.font = bold

        schedule_out = schedule_api.get_schedule(
            project_sport_id=project_sport.id,
            start_date=None,
            mobilisation_days=None,
            db=db,
            current_user=current_user,
        )
        for milestone in schedule_out.payment_schedule:
            ws.append([milestone.name, milestone.percent, milestone.date.isoformat()])
        ws.append([])

    return _xlsx_response(wb, f"{quotation.document_no}-billing-handoff.xlsx")
