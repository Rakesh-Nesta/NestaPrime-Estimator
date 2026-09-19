"""Amendment 17 (Section 23): CSV/XLSX formula-injection sanitization."""

import io

from openpyxl import load_workbook

from app.core.export_safety import sanitize_cell, sanitize_row


def test_sanitize_cell_prefixes_formula_looking_strings():
    assert sanitize_cell("=1+1") == "'=1+1"
    assert sanitize_cell("+1+1") == "'+1+1"
    assert sanitize_cell("-1+1") == "'-1+1"
    assert sanitize_cell("@SUM(A1)") == "'@SUM(A1)"


def test_sanitize_cell_leaves_ordinary_values_unchanged():
    assert sanitize_cell("Acrylic flooring") == "Acrylic flooring"
    assert sanitize_cell("PVC - Turf") == "PVC - Turf"  # doesn't start with '-'
    assert sanitize_cell(None) is None
    assert sanitize_cell(42) == 42
    assert sanitize_cell(3.14) == 3.14
    assert sanitize_cell(True) is True


def test_sanitize_row_only_touches_formula_looking_strings():
    row = ["=cmd|'/c calc'!A1", "Normal item", 5.0, None, "@evil"]
    assert sanitize_row(row) == ["'=cmd|'/c calc'!A1", "Normal item", 5.0, None, "'@evil"]


def _login(client, email, password="TestPass!1"):
    res = client.post("/auth/login", data={"username": email, "password": password})
    assert res.status_code == 200, res.text
    return {"Authorization": f"Bearer {res.json()['access_token']}"}


def _director_headers(client, director_user):
    return _login(client, "director@test.local")


def test_cost_sheet_export_neutralizes_a_formula_injection_attempt(client, director_user):
    """A malicious item_name/category/spec must never survive into the
    exported file as a live formula -- the exact CSV/XLSX-injection gap
    Amendment 17 closes. The Amount column's own live formula (set
    separately via ws.cell, not part of the sanitized row) must keep
    working unaffected."""
    headers = _director_headers(client, director_user)
    client_id = client.post(
        "/clients", json={"name": "Injection Test Client", "type": "school"}, headers=headers
    ).json()["id"]
    project = client.post(
        "/projects",
        json={
            "client_id": client_id, "city": "Mumbai", "site_condition": "level", "soil_type": "normal",
            "building_status": "open_air", "site_access": "good", "power_available": "yes",
            "water_available": True, "package": "standard",
        },
        headers=headers,
    ).json()
    cost_sheet_id = client.post(f"/projects/{project['id']}/cost-sheets", json={}, headers=headers).json()["id"]

    line = client.post(
        f"/cost-sheets/{cost_sheet_id}/lines",
        json={
            "work_package": "civil",
            "category": "=HYPERLINK(\"http://evil.example\",\"click\")",
            "item_name": "+SUM(1+1)*cmd|'/c calc'!A1",
            "spec": "@evil-spec",
            "unit": "nos",
            "quantity": 1,
            "rate": 100,
        },
        headers=headers,
    )
    assert line.status_code == 201, line.text

    res = client.get(f"/cost-sheets/{cost_sheet_id}/exports/cost-sheet", headers=headers)
    assert res.status_code == 200, res.text
    wb = load_workbook(io.BytesIO(res.content))
    ws = wb.active

    # Row 2 is the only line (E.5's structural sign-off line doesn't apply
    # to a manually-added line with no project_sport_id).
    row_values = [ws.cell(row=2, column=c).value for c in range(1, 5)]
    assert row_values[1].startswith("'"), "category must be neutralized"
    assert row_values[2].startswith("'"), "item_name must be neutralized"
    assert row_values[3].startswith("'"), "spec must be neutralized"
    assert row_values[1] == "'=HYPERLINK(\"http://evil.example\",\"click\")"
    assert row_values[2] == "'+SUM(1+1)*cmd|'/c calc'!A1"
    assert row_values[3] == "'@evil-spec"

    # The app's own live Amount formula is untouched -- not user input,
    # never sanitized.
    assert ws["H2"].value == "=F2*G2"
