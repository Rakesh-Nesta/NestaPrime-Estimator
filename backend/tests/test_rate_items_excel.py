"""Blueprint P.2 Phase 1b: "Excel rate import" (named alongside "rate
verification workflow" in the same roadmap line -- both J.1 Rate Sheet
features, distinct from Q.2 rule 6's separate Master Settings Excel
export/import, which is global %/threshold config, not per-item rates)."""

import io

from openpyxl import Workbook, load_workbook

from app.core.security import hash_password
from app.models.rate_item import RateItem
from app.models.user import User, UserRole

XLSX_CONTENT_TYPE = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"

_COLUMNS = [
    "Category", "Item name", "Spec", "Unit", "HSN/SAC", "Rate", "Vendor", "City of quote",
    "Labour category key", "Commodity watched", "Source", "Verified",
]


def _login(client, email, password="TestPass!1"):
    res = client.post("/auth/login", data={"username": email, "password": password})
    assert res.status_code == 200, res.text
    return {"Authorization": f"Bearer {res.json()['access_token']}"}


def _director_headers(client, director_user):
    return _login(client, "director@test.local")


def _procurement_headers(client, db_session):
    user = User(
        name="Test Procurement", email="procurement-rix@test.local", hashed_password=hash_password("TestPass!1"),
        role=UserRole.PROCUREMENT,
    )
    db_session.add(user)
    db_session.commit()
    return _login(client, "procurement-rix@test.local")


def _sales_headers(client, db_session):
    user = User(
        name="Test Sales", email="sales-rix@test.local", hashed_password=hash_password("TestPass!1"),
        role=UserRole.SALES,
    )
    db_session.add(user)
    db_session.commit()
    return _login(client, "sales-rix@test.local")


def _workbook_bytes(rows: list[list]) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.append(_COLUMNS)
    for row in rows:
        ws.append(row)
    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


def _import(client, headers, rows):
    content = _workbook_bytes(rows)
    return client.post(
        "/rate-items/import",
        files={"file": ("rate-sheet.xlsx", content, XLSX_CONTENT_TYPE)},
        headers=headers,
    )


def test_procurement_can_export_but_sales_cannot(client, db_session, director_user):
    procurement_headers = _procurement_headers(client, db_session)
    res = client.get("/rate-items/export", headers=procurement_headers)
    assert res.status_code == 200
    assert res.headers["content-type"] == XLSX_CONTENT_TYPE

    sales_headers = _sales_headers(client, db_session)
    assert client.get("/rate-items/export", headers=sales_headers).status_code == 403


def test_sales_cannot_import(client, db_session, director_user):
    sales_headers = _sales_headers(client, db_session)
    res = _import(client, sales_headers, [["Steel", "SHS 3x3", None, "kg", "7306", 70, None, None, None, False, "manual", False]])
    assert res.status_code == 403


def test_export_contains_a_created_rate_item(client, director_user):
    headers = _director_headers(client, director_user)
    client.post(
        "/rate-items",
        json={"category": "Steel", "item_name": "Export Probe Item", "unit": "kg", "hsn_sac": "7306", "rate": 70.0},
        headers=headers,
    )

    res = client.get("/rate-items/export", headers=headers)
    assert res.status_code == 200
    wb = load_workbook(io.BytesIO(res.content))
    ws = wb.active
    rows = list(ws.iter_rows(min_row=2, values_only=True))
    match = next(r for r in rows if r[1] == "Export Probe Item")
    assert match[0] == "Steel"  # Category
    assert match[5] == 70.0  # Rate
    assert match[10] == "manual"  # Source
    assert match[11] is False  # Verified


def test_import_creates_a_new_item_always_manual_and_unverified(client, director_user):
    """J.1: every new entry starts as a Manual, unverified rate --
    Source/Verified columns on the import file are ignored on create."""
    headers = _director_headers(client, director_user)
    res = _import(
        client, headers,
        [["Turf", "New Import Item", None, "sqm", "5703", 850.0, "Turf Co", "Mumbai", None, True, "ai", True]],
    )
    assert res.status_code == 200, res.text
    body = res.json()
    assert body["unchanged"] == 0
    assert body["errors"] == []
    assert len(body["created"]) == 1
    created = body["created"][0]
    assert created["category"] == "Turf"
    assert created["rate"] == 850.0
    assert created["vendor"] == "Turf Co"
    assert created["is_commodity_watched"] is True
    assert created["source"] == "manual"  # ignored the "ai" in the file
    assert created["verified"] is False


def test_reimporting_an_unchanged_export_is_a_safe_no_op(client, director_user):
    headers = _director_headers(client, director_user)
    client.post(
        "/rate-items",
        json={"category": "Steel", "item_name": "Noop Item", "unit": "kg", "hsn_sac": "7306", "rate": 65.0},
        headers=headers,
    )

    export_res = client.get("/rate-items/export", headers=headers)
    wb = load_workbook(io.BytesIO(export_res.content))
    ws = wb.active
    rows = [list(r) for r in ws.iter_rows(min_row=2, values_only=True)]

    import_res = client.post(
        "/rate-items/import",
        files={"file": ("rate-sheet.xlsx", export_res.content, XLSX_CONTENT_TYPE)},
        headers=headers,
    )
    assert import_res.status_code == 200, import_res.text
    body = import_res.json()
    assert body["created"] == []
    assert body["updated"] == []
    assert body["unchanged"] == len(rows)
    assert body["errors"] == []


def test_import_updates_an_existing_items_rate_through_rate_history(client, director_user, db_session):
    """A changed rate on re-import goes through the same RateHistory-
    preserving path as POST .../rate -- never a silent field overwrite."""
    headers = _director_headers(client, director_user)
    create_res = client.post(
        "/rate-items",
        json={"category": "Steel", "item_name": "History Item", "unit": "kg", "hsn_sac": "7306", "rate": 60.0},
        headers=headers,
    )
    item_id = create_res.json()["id"]

    res = _import(
        client, headers,
        [["Steel", "History Item", None, "kg", "7306", 75.0, "New Vendor", None, None, False, "manual", False]],
    )
    assert res.status_code == 200, res.text
    body = res.json()
    assert body["created"] == []
    assert len(body["updated"]) == 1
    assert body["updated"][0]["rate"] == 75.0
    assert body["updated"][0]["vendor"] == "New Vendor"

    history = client.get(f"/rate-items/{item_id}/history", headers=headers).json()
    assert len(history) == 2  # opening row (60.0, closed) + the import's new row (75.0, open)
    rates = sorted(h["rate"] for h in history)
    assert rates == [60.0, 75.0]
    open_rows = [h for h in history if h["effective_to"] is None]
    assert len(open_rows) == 1
    assert open_rows[0]["rate"] == 75.0
    assert open_rows[0]["reason"] == "Bulk Excel import"


def test_import_resolves_a_labour_category_key(client, director_user):
    headers = _director_headers(client, director_user)
    labour_categories = client.get("/labour-categories", headers=headers).json()
    ms_fab = next(c for c in labour_categories if c["key"] == "ms_fabrication_erection")

    res = _import(
        client, headers,
        [["Steel", "Labour Key Item", None, "kg", "7306", 70.0, None, None, "ms_fabrication_erection", False, "manual", False]],
    )
    assert res.status_code == 200, res.text
    created = res.json()["created"][0]
    assert created["labour_category_id"] == ms_fab["id"]


def test_import_row_with_unknown_labour_category_key_is_a_row_error_not_a_failure(client, director_user):
    headers = _director_headers(client, director_user)
    res = _import(
        client, headers,
        [
            ["Steel", "Bad Labour Key Item", None, "kg", "7306", 70.0, None, None, "not_a_real_key", False, "manual", False],
            ["Turf", "Good Item", None, "sqm", "5703", 850.0, None, None, None, False, "manual", False],
        ],
    )
    assert res.status_code == 200, res.text
    body = res.json()
    assert len(body["created"]) == 1
    assert body["created"][0]["item_name"] == "Good Item"
    assert len(body["errors"]) == 1
    assert body["errors"][0]["row"] == 2
    assert "not_a_real_key" in body["errors"][0]["detail"]


def test_import_row_missing_required_field_is_a_row_error(client, director_user):
    headers = _director_headers(client, director_user)
    res = _import(
        client, headers,
        [["", "Missing Category Item", None, "kg", "7306", 70.0, None, None, None, False, "manual", False]],
    )
    assert res.status_code == 200, res.text
    body = res.json()
    assert body["created"] == []
    assert len(body["errors"]) == 1
    assert body["errors"][0]["row"] == 2


def test_import_rejects_an_unreadable_file(client, director_user):
    headers = _director_headers(client, director_user)
    res = client.post(
        "/rate-items/import",
        files={"file": ("not-excel.xlsx", b"this is not a real xlsx file", XLSX_CONTENT_TYPE)},
        headers=headers,
    )
    assert res.status_code == 422
