import io

from openpyxl import load_workbook

from app.core.security import hash_password
from app.models.user import User, UserRole


def _login(client, email, password="TestPass!1"):
    res = client.post("/auth/login", data={"username": email, "password": password})
    assert res.status_code == 200, res.text
    return {"Authorization": f"Bearer {res.json()['access_token']}"}


def _director_headers(client, director_user):
    return _login(client, "director@test.local")


def _sales_headers(client, db_session):
    user = User(
        name="Test Sales", email="sales@test.local", hashed_password=hash_password("TestPass!1"), role=UserRole.SALES
    )
    db_session.add(user)
    db_session.commit()
    return _login(client, "sales@test.local")


def _procurement_headers(client, db_session):
    user = User(
        name="Test Procurement", email="procurement@test.local", hashed_password=hash_password("TestPass!1"),
        role=UserRole.PROCUREMENT,
    )
    db_session.add(user)
    db_session.commit()
    return _login(client, "procurement@test.local")


def _create_client_record(client, headers, client_type="government"):
    res = client.post(
        "/clients",
        json={
            "name": "Exports Client", "type": client_type,
            "contact_name": "Jane PM", "phone": "9999999999", "email": "jane@example.com",
            "billing_address": "1 Sports Lane, Mumbai",
        },
        headers=headers,
    )
    assert res.status_code == 201, res.text
    return res.json()["id"]


def _create_project(client, headers, client_id):
    fields = {
        "client_id": client_id,
        "city": "Mumbai",
        "site_condition": "level",
        "soil_type": "normal",
        "building_status": "open_air",
        "site_access": "good",
        "power_available": "yes",
        "water_available": True,
        "package": "standard",
    }
    res = client.post("/projects", json=fields, headers=headers)
    assert res.status_code == 201, res.text
    return res.json()["id"]


def _add_project_sport(client, headers, project_id, sport_key="box_cricket"):
    sport_id = next(s["id"] for s in client.get("/sports", headers=headers).json() if s["key"] == sport_key)
    res = client.post(
        f"/projects/{project_id}/sports",
        json={"sport_id": sport_id, "building_status": "open_air"},
        headers=headers,
    )
    assert res.status_code == 201, res.text
    return res.json()["id"]


def _cost_sheet_with_structure_line(client, headers):
    client_id = _create_client_record(client, headers)
    project_id = _create_project(client, headers, client_id)
    project_sport_id = _add_project_sport(client, headers, project_id)
    cost_sheet_id = client.post(f"/projects/{project_id}/cost-sheets", json={}, headers=headers).json()["id"]
    client.post(
        f"/cost-sheets/{cost_sheet_id}/structures",
        json={
            "project_sport_id": project_sport_id,
            "structure_type": "a",
            "section": "shs_3",
            "wall_thickness_mm": 2.0,
            "build_l_ft": 50,
            "build_w_ft": 25,
            "height_ft": 12,
            "steel_rate_per_kg": 68,
            "netting_rate_per_sqm": 45,
            "concrete_rate_per_cum": 6500,
        },
        headers=headers,
    )
    return project_id, project_sport_id, cost_sheet_id


def _load_xlsx(response):
    assert response.status_code == 200, response.text
    assert response.headers["content-type"] == "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    return load_workbook(io.BytesIO(response.content))


# ---------------------------------------------------------------------------
# Role gates
# ---------------------------------------------------------------------------


def test_sales_cannot_export_cost_sheet_consumption_or_bom(client, director_user, db_session):
    headers = _director_headers(client, director_user)
    _, _, cost_sheet_id = _cost_sheet_with_structure_line(client, headers)

    sales_headers = _sales_headers(client, db_session)
    for path in ("cost-sheet", "consumption-sheet", "bom", "rfq"):
        res = client.get(f"/cost-sheets/{cost_sheet_id}/exports/{path}", headers=sales_headers)
        assert res.status_code == 403, path


def test_procurement_can_export_rfq_but_not_cost_sheet(client, director_user, db_session):
    headers = _director_headers(client, director_user)
    _, _, cost_sheet_id = _cost_sheet_with_structure_line(client, headers)

    procurement_headers = _procurement_headers(client, db_session)
    rfq_res = client.get(f"/cost-sheets/{cost_sheet_id}/exports/rfq", headers=procurement_headers)
    assert rfq_res.status_code == 200

    cost_sheet_res = client.get(f"/cost-sheets/{cost_sheet_id}/exports/cost-sheet", headers=procurement_headers)
    assert cost_sheet_res.status_code == 403


# ---------------------------------------------------------------------------
# Content correctness
# ---------------------------------------------------------------------------


def test_cost_sheet_export_has_live_amount_formulas(client, director_user):
    headers = _director_headers(client, director_user)
    _, _, cost_sheet_id = _cost_sheet_with_structure_line(client, headers)

    wb = _load_xlsx(client.get(f"/cost-sheets/{cost_sheet_id}/exports/cost-sheet", headers=headers))
    ws = wb.active
    assert ws["A1"].value == "Work package"
    assert ws["H1"].value == "Amount"
    # 3 lines (steel, netting, foundation) -> rows 2-4
    assert ws["H2"].value == "=F2*G2"
    assert ws["H3"].value == "=F3*G3"
    assert ws["H4"].value == "=F4*G4"
    assert ws["B2"].value == "MS structure"


def test_consumption_sheet_export_has_live_theoretical_and_amount_formulas(client, director_user):
    headers = _director_headers(client, director_user)
    _, _, cost_sheet_id = _cost_sheet_with_structure_line(client, headers)

    wb = _load_xlsx(client.get(f"/cost-sheets/{cost_sheet_id}/exports/consumption-sheet", headers=headers))
    ws = wb.active
    assert ws["A1"].value == "Category"
    assert ws["D1"].value == "Theoretical qty"
    assert ws["D2"].value == "=F2/(1+E2/100)"
    assert ws["H2"].value == "=F2*G2"
    assert ws["E2"].value == 5.0  # steel line wastage %
    assert ws["I2"].value is None  # Vendor -- blank, no PO raised yet
    assert ws["J2"].value is None  # PO No. -- blank
    assert ws["K2"].value is None  # Delivery date -- blank
    assert ws["L2"].value is None  # Received qty -- blank
    assert ws["M2"].value is None  # Balance -- blank


def test_consumption_sheet_export_shows_real_po_data_once_raised(client, director_user):
    """Part O: once a PO is raised for a line, the export reads it live --
    not fabricated, not left permanently blank."""
    headers = _director_headers(client, director_user)
    _, _, cost_sheet_id = _cost_sheet_with_structure_line(client, headers)
    line_id = client.get(f"/cost-sheets/{cost_sheet_id}/lines", headers=headers).json()[0]["id"]
    vendor_id = client.post("/vendors", json={"name": "Steel Co"}, headers=headers).json()["id"]
    client.post(
        f"/cost-sheets/{cost_sheet_id}/purchase-orders",
        json={
            "vendor_id": vendor_id,
            "lines": [{"cost_sheet_line_id": line_id, "quantity": 100, "rate": 70}],
            "delivery_date": "2026-10-01",
        },
        headers=headers,
    )

    wb = _load_xlsx(client.get(f"/cost-sheets/{cost_sheet_id}/exports/consumption-sheet", headers=headers))
    ws = wb.active
    assert ws["I2"].value == "Steel Co"
    assert ws["J2"].value.startswith("PO-")
    assert ws["K2"].value == "2026-10-01"
    assert ws["L2"].value == 0.0
    assert ws["M2"].value == 100.0


def test_rfq_export_has_only_item_spec_unit_qty(client, director_user):
    """J.4: 'item, spec, qty, unit only; no rates, no vendor names.'"""
    headers = _director_headers(client, director_user)
    _, _, cost_sheet_id = _cost_sheet_with_structure_line(client, headers)

    wb = _load_xlsx(client.get(f"/cost-sheets/{cost_sheet_id}/exports/rfq", headers=headers))
    ws = wb.active
    assert [c.value for c in ws[1]] == ["Item", "Spec", "Unit", "Qty"]
    assert ws.max_column == 4
    body_rows = [tuple(c.value for c in row) for row in ws.iter_rows(min_row=2)]
    assert all(len(row) == 4 for row in body_rows)
    assert any("SHS 3" in str(row[0]) for row in body_rows)


def test_bom_export_is_sorted_by_category_and_includes_source(client, director_user):
    headers = _director_headers(client, director_user)
    _, _, cost_sheet_id = _cost_sheet_with_structure_line(client, headers)

    wb = _load_xlsx(client.get(f"/cost-sheets/{cost_sheet_id}/exports/bom", headers=headers))
    ws = wb.active
    assert ws["A1"].value == "Category"
    assert ws["H1"].value == "Source"
    categories = [ws.cell(row=r, column=1).value for r in range(2, ws.max_row + 1)]
    assert categories == sorted(categories)
    sources = [ws.cell(row=r, column=8).value for r in range(2, ws.max_row + 1)]
    assert all(s == "manual" for s in sources)


# ---------------------------------------------------------------------------
# Billing Handoff
# ---------------------------------------------------------------------------


def _won_quotation(client, headers):
    client_id = _create_client_record(client, headers, client_type="government")
    project_id = _create_project(client, headers, client_id)
    project_sport_id = _add_project_sport(client, headers, project_id)
    cost_sheet_id = client.post(
        f"/projects/{project_id}/cost-sheets", json={"cost_total": 850000}, headers=headers
    ).json()["id"]
    client.post(f"/cost-sheets/{cost_sheet_id}/verify", headers=headers)

    estimate = client.post(
        f"/projects/{project_id}/estimates",
        json={"options": [{"project_sport_id": project_sport_id, "package": "standard", "cost_for_option": 850000}]},
        headers=headers,
    ).json()
    option_id = estimate["options"][0]["id"]
    client.patch(
        f"/estimates/{estimate['id']}/options/{option_id}/client-status",
        json={"client_status": "approved", "waive_evidence_reason": "test setup"},
        headers=headers,
    )
    quotation = client.post(
        f"/projects/{project_id}/quotations",
        json={"estimate_id": estimate["id"], "included_option_ids": [option_id]},
        headers=headers,
    ).json()
    client.post(f"/quotations/{quotation['id']}/release", headers=headers)
    client.post(f"/quotations/{quotation['id']}/send", headers=headers)
    won = client.post(
        f"/quotations/{quotation['id']}/mark-won",
        json={"reason": "Client accepted", "waive_evidence_reason": "test setup"},
        headers=headers,
    )
    assert won.status_code == 200, won.text
    return quotation["id"], won.json()


def test_billing_handoff_requires_won_status(client, director_user):
    headers = _director_headers(client, director_user)
    client_id = _create_client_record(client, headers)
    project_id = _create_project(client, headers, client_id)
    project_sport_id = _add_project_sport(client, headers, project_id)
    cost_sheet_id = client.post(
        f"/projects/{project_id}/cost-sheets", json={"cost_total": 850000}, headers=headers
    ).json()["id"]
    client.post(f"/cost-sheets/{cost_sheet_id}/verify", headers=headers)
    estimate = client.post(
        f"/projects/{project_id}/estimates",
        json={"options": [{"project_sport_id": project_sport_id, "package": "standard", "cost_for_option": 850000}]},
        headers=headers,
    ).json()
    option_id = estimate["options"][0]["id"]
    client.patch(
        f"/estimates/{estimate['id']}/options/{option_id}/client-status",
        json={"client_status": "approved", "waive_evidence_reason": "test setup"},
        headers=headers,
    )
    quotation = client.post(
        f"/projects/{project_id}/quotations",
        json={"estimate_id": estimate["id"], "included_option_ids": [option_id]},
        headers=headers,
    ).json()

    res = client.get(f"/quotations/{quotation['id']}/exports/billing-handoff", headers=headers)
    assert res.status_code == 400


def test_billing_handoff_carries_client_total_and_payment_schedule_no_cost(client, director_user, db_session):
    """J.4: 'client name, contact, address, quotation number and date, the
    payment schedule/milestones, and the total -- nothing else, no cost,
    no margin, no HSN.' Also confirms Sales can use it (K.3-safe by
    construction)."""
    headers = _director_headers(client, director_user)
    quotation_id, quotation = _won_quotation(client, headers)

    sales_headers = _sales_headers(client, db_session)
    wb = _load_xlsx(client.get(f"/quotations/{quotation_id}/exports/billing-handoff", headers=sales_headers))
    ws = wb.active

    values = {ws.cell(row=r, column=1).value: ws.cell(row=r, column=2).value for r in range(1, 9)}
    assert values["Client"] == "Exports Client"
    assert values["Contact"] == "Jane PM"
    assert values["Billing address"] == "1 Sports Lane, Mumbai"
    assert values["Quotation No."] == quotation["document_no"]
    assert values["Total (GST-inclusive, 18% flat)"] == quotation["quotation_total"]

    all_text = " ".join(str(ws.cell(row=r, column=c).value) for r in range(1, ws.max_row + 1) for c in (1, 2, 3))
    assert "cost" not in all_text.lower().replace("gst-inclusive", "")
    assert "margin" not in all_text.lower()
    assert "hsn" not in all_text.lower()
    assert "Payment schedule" in all_text
