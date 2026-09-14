"""Amendment 11 Part A (Section 10): a RateItem can be created "awaiting
rate" (rate=None) -- a real, HSN/SAC-classified catalog entry with no
defensible Rs/unit figure yet, never a fabricated number. Cannot be
confirmed into an AI rate until a PM/Director enters a real value via
POST /rate-items/{id}/rate."""

import io

from openpyxl import Workbook, load_workbook

from app.models.rate_item import RateItem, RateSource
from app.models.rate_history import RateHistory


def _login(client, director_user):
    res = client.post(
        "/auth/login",
        data={"username": "director@test.local", "password": "TestPass!1"},
    )
    return {"Authorization": f"Bearer {res.json()['access_token']}"}


AWAITING_RATE_FIELDS = {
    "category": "Civil",
    "item_name": "Asphalt base course",
    "unit": "sqmt",
    "hsn_sac": "995454",
    # rate omitted entirely -- defaults to None per RateItemCreate
}


def test_rate_item_can_be_created_with_no_rate(client, director_user):
    headers = _login(client, director_user)
    res = client.post("/rate-items", json=AWAITING_RATE_FIELDS, headers=headers)
    assert res.status_code == 201, res.text
    body = res.json()
    assert body["rate"] is None
    assert body["source"] == "manual"
    assert body["verified"] is False


def test_creating_an_awaiting_rate_item_writes_no_rate_history(client, director_user, db_session):
    headers = _login(client, director_user)
    res = client.post("/rate-items", json=AWAITING_RATE_FIELDS, headers=headers)
    item_id = res.json()["id"]

    history = db_session.query(RateHistory).filter(RateHistory.rate_item_id == item_id).all()
    assert history == []


def test_cannot_confirm_an_awaiting_rate_item(client, director_user):
    headers = _login(client, director_user)
    create_res = client.post("/rate-items", json=AWAITING_RATE_FIELDS, headers=headers)
    item_id = create_res.json()["id"]

    res = client.post(f"/rate-items/{item_id}/confirm", headers=headers)
    assert res.status_code == 400
    assert "awaiting-rate" in res.json()["detail"].lower()


def test_setting_the_first_real_rate_on_an_awaiting_rate_item(client, director_user, db_session):
    """POST .../rate is how an awaiting-rate item gets its first real,
    quantity-backed number -- must not crash comparing against a None
    previous rate, and must write exactly one open RATE_HISTORY row."""
    headers = _login(client, director_user)
    create_res = client.post("/rate-items", json=AWAITING_RATE_FIELDS, headers=headers)
    item_id = create_res.json()["id"]

    res = client.post(
        f"/rate-items/{item_id}/rate",
        json={"rate": 450.0, "reason": "First real rate from accountant sign-off"},
        headers=headers,
    )
    assert res.status_code == 200, res.text
    body = res.json()["rate_item"]
    assert body["rate"] == 450.0
    assert res.json()["commodity_alert"] is None

    history = db_session.query(RateHistory).filter(RateHistory.rate_item_id == item_id).all()
    assert len(history) == 1
    assert history[0].effective_to is None
    assert float(history[0].rate) == 450.0


def test_can_confirm_once_a_real_rate_is_set(client, director_user):
    headers = _login(client, director_user)
    create_res = client.post("/rate-items", json=AWAITING_RATE_FIELDS, headers=headers)
    item_id = create_res.json()["id"]
    client.post(f"/rate-items/{item_id}/rate", json={"rate": 450.0, "reason": "test"}, headers=headers)

    res = client.post(f"/rate-items/{item_id}/confirm", headers=headers)
    assert res.status_code == 200, res.text
    assert res.json()["source"] == "ai"


def test_bulk_category_rate_update_skips_awaiting_rate_items(client, director_user):
    headers = _login(client, director_user)
    aw = client.post("/rate-items", json=AWAITING_RATE_FIELDS, headers=headers).json()
    rated = client.post(
        "/rate-items",
        json={**AWAITING_RATE_FIELDS, "item_name": "Rated civil item", "rate": 100.0},
        headers=headers,
    ).json()

    res = client.post(
        "/rate-items/bulk-rate-update",
        json={"category": "Civil", "percent_change": 10.0, "reason": "test bulk move"},
        headers=headers,
    )
    assert res.status_code == 200, res.text
    updated_ids = {i["rate_item"]["id"] for i in res.json()["items"]}
    assert aw["id"] not in updated_ids
    assert rated["id"] in updated_ids


def test_sync_draft_lines_rejects_an_awaiting_rate_item(client, director_user):
    headers = _login(client, director_user)
    create_res = client.post("/rate-items", json=AWAITING_RATE_FIELDS, headers=headers)
    item_id = create_res.json()["id"]

    res = client.post(f"/rate-items/{item_id}/sync-draft-lines", headers=headers)
    assert res.status_code == 400
    assert "awaiting-rate" in res.json()["detail"].lower()


def test_export_includes_a_blank_cell_for_an_awaiting_rate_item(client, director_user):
    headers = _login(client, director_user)
    client.post(
        "/rate-items", json={**AWAITING_RATE_FIELDS, "item_name": "Export Awaiting Test"}, headers=headers
    )

    res = client.get("/rate-items/export", headers=headers)
    assert res.status_code == 200, res.text
    wb = load_workbook(io.BytesIO(res.content))
    ws = wb.active
    row = next(r for r in ws.iter_rows(min_row=2, values_only=True) if r[1] == "Export Awaiting Test")
    assert row[5] is None  # Rate column


def test_import_creates_an_awaiting_rate_item_from_a_blank_rate_cell(client, director_user, db_session):
    """A blank Rate cell on a new row is 'awaiting rate', not an error."""
    headers = _login(client, director_user)
    wb = Workbook()
    ws = wb.active
    ws.append(
        ["Category", "Item name", "Spec", "Unit", "HSN/SAC", "Rate", "Vendor", "City of quote",
         "Labour category key", "Commodity watched", "Source", "Verified"]
    )
    ws.append(["Civil", "Imported Awaiting Item", None, "sqmt", "995454", None, None, None, None, False, None, None])
    buffer = io.BytesIO()
    wb.save(buffer)

    res = client.post(
        "/rate-items/import",
        files={"file": ("rate-sheet.xlsx", buffer.getvalue(),
                         "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
        headers=headers,
    )
    assert res.status_code == 200, res.text
    body = res.json()
    assert body["errors"] == []
    assert len(body["created"]) == 1
    assert body["created"][0]["rate"] is None

    history = (
        db_session.query(RateHistory)
        .filter(RateHistory.rate_item_id == body["created"][0]["id"])
        .all()
    )
    assert history == []


def test_reimporting_an_unchanged_awaiting_rate_export_is_a_no_op(client, director_user):
    """Documented invariant (export_rate_items docstring): export-then-
    reimport unchanged is always a safe no-op, including for an
    awaiting-rate item's blank Rate cell."""
    headers = _login(client, director_user)
    client.post(
        "/rate-items", json={**AWAITING_RATE_FIELDS, "item_name": "Roundtrip Awaiting Item"}, headers=headers
    )

    export_res = client.get("/rate-items/export", headers=headers)
    import_res = client.post(
        "/rate-items/import",
        files={
            "file": (
                "rate-sheet.xlsx", export_res.content,
                "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            )
        },
        headers=headers,
    )
    assert import_res.status_code == 200, import_res.text
    body = import_res.json()
    assert body["errors"] == []
    assert body["created"] == []
    assert len(body["updated"]) == 0
