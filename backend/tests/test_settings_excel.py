import io
from datetime import date

from openpyxl import Workbook, load_workbook

from app.core.security import hash_password
from app.models.user import User, UserRole

XLSX_CONTENT_TYPE = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


def _login(client, email, password="TestPass!1"):
    res = client.post("/auth/login", data={"username": email, "password": password})
    assert res.status_code == 200, res.text
    return {"Authorization": f"Bearer {res.json()['access_token']}"}


def _director_headers(client, director_user):
    return _login(client, "director@test.local")


def _pm_headers(client, db_session):
    user = User(name="Test PM", email="pm@test.local", hashed_password=hash_password("TestPass!1"), role=UserRole.PM)
    db_session.add(user)
    db_session.commit()
    return _login(client, "pm@test.local")


def _sales_headers(client, db_session):
    user = User(
        name="Test Sales", email="sales@test.local", hashed_password=hash_password("TestPass!1"), role=UserRole.SALES
    )
    db_session.add(user)
    db_session.commit()
    return _login(client, "sales@test.local")


def _workbook_bytes(rows: list[list]) -> bytes:
    """rows: each a [key, scope, scope_value, value, unit, effective_from, reason]
    list, matching export_settings()'s own column order exactly."""
    wb = Workbook()
    ws = wb.active
    ws.append(["Key", "Scope", "Scope value", "Value", "Unit", "Effective from", "Reason"])
    for row in rows:
        ws.append(row)
    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


def _import(client, headers, rows):
    content = _workbook_bytes(rows)
    return client.post(
        "/settings/import",
        files={"file": ("settings.xlsx", content, XLSX_CONTENT_TYPE)},
        headers=headers,
    )


def test_pm_can_export_but_sales_cannot(client, db_session, director_user):
    pm_headers = _pm_headers(client, db_session)
    res = client.get("/settings/export", headers=pm_headers)
    assert res.status_code == 200
    assert res.headers["content-type"] == XLSX_CONTENT_TYPE

    sales_headers = _sales_headers(client, db_session)
    assert client.get("/settings/export", headers=sales_headers).status_code == 403


def test_export_contains_current_settings(client, director_user):
    headers = _director_headers(client, director_user)
    client.post("/settings", json={"key": "export_probe_key", "value": "42", "unit": "%", "reason": "seed"}, headers=headers)

    res = client.get("/settings/export", headers=headers)
    assert res.status_code == 200
    wb = load_workbook(io.BytesIO(res.content))
    ws = wb.active
    rows = list(ws.iter_rows(min_row=2, values_only=True))
    match = next(r for r in rows if r[0] == "export_probe_key")
    assert match[3] == "42"  # Value column
    assert match[4] == "%"  # Unit column


def test_pm_cannot_import(client, director_user, db_session):
    pm_headers = _pm_headers(client, db_session)
    res = _import(client, pm_headers, [["gst_rate_percent", "global", None, "20.0", None, None, "pm attempt"]])
    assert res.status_code == 403


def test_import_creates_a_new_version_for_a_changed_value(client, director_user):
    headers = _director_headers(client, director_user)
    client.post("/settings", json={"key": "steel_rate", "value": "68", "reason": "baseline"}, headers=headers)

    res = _import(client, headers, [["steel_rate", "global", None, "72", None, None, "Excel edit"]])
    assert res.status_code == 200, res.text
    body = res.json()
    assert body["unchanged"] == 0
    assert len(body["created"]) == 1
    assert body["created"][0]["key"] == "steel_rate"
    assert body["created"][0]["value"] == "72"
    assert body["created"][0]["reason"] == "Excel edit"
    assert body["errors"] == []

    current = next(s for s in client.get("/settings", headers=headers).json() if s["key"] == "steel_rate")
    assert current["value"] == "72"
    history = client.get("/settings/steel_rate/history", headers=headers).json()
    assert len(history) == 2  # Q.2 rule 1 -- a new version, not a mutation


def test_import_skips_rows_with_an_unchanged_value(client, director_user):
    headers = _director_headers(client, director_user)
    client.post("/settings", json={"key": "unchanged_rate", "value": "100", "reason": "baseline"}, headers=headers)

    res = _import(client, headers, [["unchanged_rate", "global", None, "100", None, None, None]])
    assert res.status_code == 200, res.text
    body = res.json()
    assert body["unchanged"] == 1
    assert body["created"] == []

    history = client.get("/settings/unchanged_rate/history", headers=headers).json()
    assert len(history) == 1  # no pointless new version


def test_export_then_reimport_unchanged_is_a_safe_no_op(client, director_user):
    headers = _director_headers(client, director_user)
    client.post("/settings", json={"key": "roundtrip_a", "value": "5", "reason": "seed"}, headers=headers)
    client.post("/settings", json={"key": "roundtrip_b", "value": "10", "unit": "%", "reason": "seed"}, headers=headers)

    exported = client.get("/settings/export", headers=headers).content
    res = client.post(
        "/settings/import", files={"file": ("export.xlsx", exported, XLSX_CONTENT_TYPE)}, headers=headers
    )
    assert res.status_code == 200, res.text
    body = res.json()
    assert body["created"] == []
    assert body["errors"] == []
    # every currently-effective setting round-tripped as "unchanged"
    current_count = len(client.get("/settings", headers=headers).json())
    assert body["unchanged"] == current_count


def test_import_defaults_reason_and_effective_from_when_blank(client, director_user):
    headers = _director_headers(client, director_user)
    res = _import(client, headers, [["blank_defaults_key", None, None, "1", None, None, None]])
    assert res.status_code == 200, res.text
    created = res.json()["created"][0]
    assert created["reason"] == "Bulk Excel import"
    assert created["effective_from"] == date.today().isoformat()
    assert created["scope"] == "global"


def test_import_reports_a_bad_row_without_aborting_the_rest(client, director_user):
    headers = _director_headers(client, director_user)
    res = _import(
        client,
        headers,
        [
            [None, "global", None, "1", None, None, None],  # missing key -> row error
            ["good_row_key", "not_a_real_scope", None, "1", None, None, None],  # bad scope -> row error
            ["good_row_key_2", "global", None, "1", None, None, None],  # valid
        ],
    )
    assert res.status_code == 200, res.text
    body = res.json()
    assert len(body["errors"]) == 2
    assert {e["row"] for e in body["errors"]} == {2, 3}
    assert len(body["created"]) == 1
    assert body["created"][0]["key"] == "good_row_key_2"


def test_import_rejects_a_non_excel_file(client, director_user):
    headers = _director_headers(client, director_user)
    res = client.post(
        "/settings/import",
        files={"file": ("not_a_workbook.xlsx", b"this is not a real xlsx file", XLSX_CONTENT_TYPE)},
        headers=headers,
    )
    assert res.status_code == 422


def test_import_respects_city_scope_and_scope_value(client, director_user):
    """A city-scoped row (e.g. a per-city regional constant) must import
    to that specific scope_value, not collide with the global row for
    the same key."""
    headers = _director_headers(client, director_user)
    client.post(
        "/settings",
        json={"key": "city_scoped_key", "scope": "city", "scope_value": "Mumbai", "value": "1.1", "reason": "seed"},
        headers=headers,
    )

    res = _import(
        client, headers, [["city_scoped_key", "city", "Mumbai", "1.25", None, None, "correction"]]
    )
    assert res.status_code == 200, res.text
    assert len(res.json()["created"]) == 1

    history = client.get(
        "/settings/city_scoped_key/history", params={"scope_value": "Mumbai"}, headers=headers
    ).json()
    assert len(history) == 2
    assert history[0]["value"] == "1.25"


def test_settings_export_import_require_auth(client):
    assert client.get("/settings/export").status_code == 401
    assert client.post("/settings/import", files={"file": ("x.xlsx", b"", XLSX_CONTENT_TYPE)}).status_code == 401
